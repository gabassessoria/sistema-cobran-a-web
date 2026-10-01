
from __future__ import annotations

from datetime import datetime, date, timedelta
from io import BytesIO
import re
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


SHEET_NAME = "DEMANDAS"

COL_SEI = "NUMERO SEI"
COL_DATA = "DATA  DO RECEBIMENTO DA DEMANDA"
COL_PRAZO = "PRAZO "
COL_STATUS = "ANDAMENTO"
COL_INTERESSADO = "INTERESSADO"
COL_AREA = "ÁREA TÉCNICA RESPONSÁVEL PELA NOTA (SX-GXX-DXX)"
COL_PRIORIDADE = "PRIORITÁRIOS"

OBRIGATORIAS = [
    COL_SEI,
    COL_DATA,
    COL_PRAZO,
    COL_STATUS,
    COL_INTERESSADO,
    COL_AREA,
    COL_PRIORIDADE,
]


def _normalizar_texto(serie: pd.Series) -> pd.Series:
    return serie.fillna("").astype(str).str.strip()


def ler_demandas(conteudo: bytes) -> pd.DataFrame:
    df = pd.read_excel(BytesIO(conteudo), sheet_name=SHEET_NAME)
    faltando = [c for c in OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(faltando))

    df = df[
        df[COL_SEI].notna()
        & (_normalizar_texto(df[COL_SEI]) != "")
    ].copy()
    return df


def opcoes_do_arquivo(conteudo: bytes) -> dict:
    df = ler_demandas(conteudo)

    status = sorted({
        s for s in _normalizar_texto(df[COL_STATUS]).tolist()
        if s
    })

    interessados = {}
    for prioridade in ("Todos", "SIM", "NÃO"):
        base = df
        if prioridade in ("SIM", "NÃO"):
            base = base[
                _normalizar_texto(base[COL_PRIORIDADE]).str.upper().eq(prioridade)
            ]
        valores = sorted({
            x for x in _normalizar_texto(base[COL_INTERESSADO]).tolist()
            if x
        })
        interessados[prioridade] = valores

    return {
        "status": status,
        "interessados": interessados,
        "total": int(len(df)),
    }


def _extrair_prazo_valor(valor):
    if pd.isna(valor):
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if valor > 0 else None
    try:
        num = float(str(valor).strip())
        return num if num > 0 else None
    except Exception:
        return None


def _calcular_data_vencimento(row):
    data_recebimento = row.get("DATA_RECEBIMENTO")
    prazo = row.get("PRAZO_VALOR")

    prazo_str = str(row.get(COL_PRAZO, "")).strip() if pd.notna(row.get(COL_PRAZO)) else ""
    if "PRAZO ADMINISTRATIVO" in prazo_str.upper() or prazo_str == "-":
        return "PRAZO ADMINISTRATIVO"

    if pd.notna(prazo) and isinstance(prazo, (int, float)) and prazo > 0:
        if pd.notna(data_recebimento) and isinstance(data_recebimento, (datetime, date, pd.Timestamp)):
            return data_recebimento + timedelta(days=int(prazo))
    return None


def preparar_datas(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["DATA_RECEBIMENTO"] = pd.to_datetime(df[COL_DATA], errors="coerce")
    df["PRAZO_VALOR"] = df[COL_PRAZO].apply(_extrair_prazo_valor)
    df["DATA_VENCIMENTO"] = df.apply(_calcular_data_vencimento, axis=1)
    return df


def filtrar_prazos(df: pd.DataFrame, tipos: list[str], hoje: date | None = None) -> pd.DataFrame:
    hoje = hoje or datetime.now().date()
    tipos = [x for x in tipos if x]

    if not tipos or "todos" in tipos:
        return df.copy()

    administrativo = (
        df["DATA_VENCIMENTO"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
        .eq("PRAZO ADMINISTRATIVO")
    )
    datas = pd.to_datetime(
        df["DATA_VENCIMENTO"].where(~administrativo),
        errors="coerce"
    ).dt.date

    mascara = pd.Series(False, index=df.index)

    if "vencidos" in tipos:
        mascara |= (~administrativo) & datas.notna() & (datas < hoje)

    if "a_vencer" in tipos:
        limite = hoje + timedelta(days=5)
        mascara |= (
            (~administrativo)
            & datas.notna()
            & (datas >= hoje)
            & (datas <= limite)
        )

    if "prazo_administrativo" in tipos:
        mascara |= administrativo

    return df[mascara].copy()


def _safe_filename(texto: str) -> str:
    texto = re.sub(r'[<>:"/\\|?*]+', "_", texto or "")
    texto = re.sub(r"\s+", "_", texto.strip())
    return texto[:60] or "relatorio"


def _ajustar_larguras(ws, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_length = 0
        for cell in col:
            valor = cell.value
            if isinstance(valor, (datetime, date)):
                tamanho = 10
            else:
                tamanho = len(str(valor or ""))
            max_length = max(max_length, tamanho)
        ws.column_dimensions[col_letter].width = min(max(max_length + 2, 12), max_width)


def gerar_relatorio(
    conteudo: bytes,
    prazos: list[str],
    prioridade: str,
    status: list[str],
    interessados: list[str],
) -> tuple[bytes, str, dict]:
    df = preparar_datas(ler_demandas(conteudo))

    prioridade = (prioridade or "Todos").strip()
    if prioridade in ("SIM", "NÃO"):
        df = df[
            _normalizar_texto(df[COL_PRIORIDADE]).str.upper().eq(prioridade)
        ]

    if status:
        df = df[
            _normalizar_texto(df[COL_STATUS]).isin(status)
        ]

    if interessados:
        df = df[
            _normalizar_texto(df[COL_INTERESSADO]).isin(interessados)
        ]

    df = filtrar_prazos(df, prazos)
    hoje = datetime.now().date()

    saida = df[
        [COL_SEI, COL_INTERESSADO, COL_AREA, COL_PRIORIDADE, COL_STATUS, "DATA_VENCIMENTO"]
    ].copy()
    saida.columns = [
        "NÚMERO SEI",
        "INTERESSADO",
        "ÁREA TÉCNICA",
        "PRIORITÁRIO",
        "ANDAMENTO",
        "DATA DE VENCIMENTO",
    ]

    def _ordem(valor):
        if valor is None or valor == "PRAZO ADMINISTRATIVO" or (isinstance(valor, float) and pd.isna(valor)):
            return pd.Timestamp.max
        try:
            return pd.Timestamp(valor)
        except Exception:
            return pd.Timestamp.max

    if not saida.empty:
        saida["_ORDEM"] = saida["DATA DE VENCIMENTO"].apply(_ordem)
        saida = saida.sort_values(["_ORDEM", "NÚMERO SEI"], na_position="last").drop(columns=["_ORDEM"])

    output = BytesIO()
    saida.to_excel(output, index=False)
    output.seek(0)

    wb = load_workbook(output)
    ws = wb.active
    ws.title = "RELATÓRIO"

    azul = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    vermelho = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    amarelo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    cinza = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    branco = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = azul
        cell.font = branco
        cell.alignment = Alignment(horizontal="center")

    cabecalhos = {cell.value: cell.column for cell in ws[1]}
    col_venc = cabecalhos.get("DATA DE VENCIMENTO")

    if col_venc:
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, col_venc)
            valor = cell.value
            fill = None
            fonte = None
            if isinstance(valor, (datetime, date)):
                cell.number_format = "dd/mm/yyyy"
                data_v = valor.date() if isinstance(valor, datetime) else valor
                dias = (data_v - hoje).days
                if dias < 0:
                    fill = vermelho
                    fonte = branco
                elif 0 <= dias <= 5:
                    fill = amarelo
            elif valor == "PRAZO ADMINISTRATIVO":
                fill = cinza

            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row_idx, c).fill = fill
                    if fonte:
                        ws.cell(row_idx, c).font = fonte

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    _ajustar_larguras(ws)

    final = BytesIO()
    wb.save(final)

    tipos = "todos" if not prazos or "todos" in prazos else "_".join(prazos)
    nome = f"processos_{_safe_filename(tipos)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    resumo = {
        "total": int(len(saida)),
        "prioridade": prioridade,
    }
    return final.getvalue(), nome, resumo


def identificar_unidade_acompanhamento(area_tecnica):
    if pd.isna(area_tecnica):
        return None

    texto = str(area_tecnica).strip().upper()
    if not texto:
        return None

    match_cd = re.match(r"^(CD-DIRETOR-\d+)(?=[\s\-(]|$)", texto)
    if match_cd:
        return match_cd.group(1)

    superintendencias = {"SA", "SD", "SF", "SH", "SR"}
    match = re.match(r"^([A-Z]+)(?=[\s\-(]|$)", texto)
    if not match:
        return None

    sigla = match.group(1)
    return sigla if sigla in superintendencias else None


def _xlsx_unidade(df_unidade: pd.DataFrame) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    azul = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    vermelho = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    amarelo = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    cinza = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    branco = Font(color="FFFFFF", bold=True)

    hoje = datetime.now().date()
    cabecalhos = ["NÚMERO SEI", "INTERESSADO", "ÁREA TÉCNICA", "VENCIMENTO"]

    for prioridade, nome_aba in (("SIM", "PRIORITÁRIOS"), ("NÃO", "NÃO PRIORITÁRIOS")):
        ws = wb.create_sheet(nome_aba)

        for idx, cab in enumerate(cabecalhos, 1):
            c = ws.cell(1, idx, cab)
            c.fill = azul
            c.font = branco
            c.alignment = Alignment(horizontal="center")

        base = df_unidade[df_unidade["PRIORIDADE_NORMALIZADA"] == prioridade].copy()

        def _ordem(valor):
            if valor == "PRAZO ADMINISTRATIVO" or pd.isna(valor):
                return pd.Timestamp.max
            try:
                return pd.Timestamp(valor)
            except Exception:
                return pd.Timestamp.max

        if not base.empty:
            base["_ORDEM"] = base["DATA_VENCIMENTO"].apply(_ordem)
            base = base.sort_values(["_ORDEM", COL_SEI], na_position="last")

        for linha, (_, row) in enumerate(base.iterrows(), start=2):
            valores = [
                row[COL_SEI],
                row[COL_INTERESSADO],
                row[COL_AREA],
                row["DATA_VENCIMENTO"],
            ]
            for col, valor in enumerate(valores, start=1):
                c = ws.cell(linha, col, valor)
                if col == 4 and isinstance(valor, (datetime, date, pd.Timestamp)):
                    c.number_format = "dd/mm/yyyy"

            venc = row["DATA_VENCIMENTO"]
            fill = None
            fonte = None

            if venc == "PRAZO ADMINISTRATIVO":
                fill = cinza
            elif prioridade == "SIM" and isinstance(venc, (datetime, date, pd.Timestamp)):
                data_v = venc.date() if isinstance(venc, (datetime, pd.Timestamp)) else venc
                dias = (data_v - hoje).days
                if dias < 0:
                    fill = vermelho
                    fonte = branco
                elif 0 <= dias <= 5:
                    fill = amarelo

            if fill:
                for col in range(1, 5):
                    ws.cell(linha, col).fill = fill
                    if fonte:
                        ws.cell(linha, col).font = fonte

        ws.auto_filter.ref = f"A1:D{max(ws.max_row, 1)}"
        ws.freeze_panes = "A2"
        _ajustar_larguras(ws, max_width=45)

    buff = BytesIO()
    wb.save(buff)
    return buff.getvalue()


def _xlsx_fora(df_fora: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FORA DA DISTRIBUIÇÃO"

    azul = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    branco = Font(color="FFFFFF", bold=True)

    cabecalhos = ["NÚMERO SEI", "INTERESSADO", "ÁREA TÉCNICA", "PRIORITÁRIO", "VENCIMENTO"]
    for idx, cab in enumerate(cabecalhos, 1):
        c = ws.cell(1, idx, cab)
        c.fill = azul
        c.font = branco
        c.alignment = Alignment(horizontal="center")

    base = df_fora.sort_values([COL_AREA, COL_SEI], na_position="last") if not df_fora.empty else df_fora

    for linha, (_, row) in enumerate(base.iterrows(), start=2):
        valores = [
            row[COL_SEI],
            row[COL_INTERESSADO],
            row[COL_AREA],
            row["PRIORIDADE_NORMALIZADA"],
            row["DATA_VENCIMENTO"],
        ]
        for col, valor in enumerate(valores, 1):
            c = ws.cell(linha, col, valor)
            if col == 5 and isinstance(valor, (datetime, date, pd.Timestamp)):
                c.number_format = "dd/mm/yyyy"

    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
    ws.freeze_panes = "A2"
    _ajustar_larguras(ws, max_width=55)

    buff = BytesIO()
    wb.save(buff)
    return buff.getvalue()


def analisar_acompanhamento(conteudo: bytes) -> dict:
    df = preparar_datas(ler_demandas(conteudo))

    df = df[
        _normalizar_texto(df[COL_STATUS]).str.upper().eq("ANÁLISE TÉCNICA")
    ].copy()

    problemas = []
    for coluna, nome in (
        (COL_INTERESSADO, "INTERESSADO"),
        (COL_AREA, "ÁREA TÉCNICA"),
        (COL_PRIORIDADE, "PRIORITÁRIOS"),
        (COL_DATA, "DATA DO RECEBIMENTO"),
        (COL_PRAZO, "PRAZO"),
    ):
        vazio = df[coluna].isna() | _normalizar_texto(df[coluna]).eq("")
        qtd = int(vazio.sum())
        if qtd:
            problemas.append(f"{qtd} processo(s) sem {nome}")

    df["PRIORIDADE_NORMALIZADA"] = _normalizar_texto(df[COL_PRIORIDADE]).str.upper()
    invalidos = ~df["PRIORIDADE_NORMALIZADA"].isin(["SIM", "NÃO"])
    if invalidos.any():
        problemas.append(f"{int(invalidos.sum())} processo(s) com PRIORITÁRIOS diferente de SIM/NÃO")

    df["UNIDADE_ACOMPANHAMENTO"] = df[COL_AREA].apply(identificar_unidade_acompanhamento)
    fora = df["UNIDADE_ACOMPANHAMENTO"].isna()

    aptos = df[
        df["UNIDADE_ACOMPANHAMENTO"].notna()
        & df["PRIORIDADE_NORMALIZADA"].isin(["SIM", "NÃO"])
    ].copy()

    unidades = ["SA", "SD", "SF", "SH", "SR"]
    conselhos = sorted(
        u for u in aptos["UNIDADE_ACOMPANHAMENTO"].dropna().unique()
        if str(u).startswith("CD-DIRETOR-")
    )
    unidades.extend(conselhos)

    resumo = []
    for unidade in unidades:
        bloco = aptos[aptos["UNIDADE_ACOMPANHAMENTO"] == unidade]
        if bloco.empty:
            continue
        p = int((bloco["PRIORIDADE_NORMALIZADA"] == "SIM").sum())
        n = int((bloco["PRIORIDADE_NORMALIZADA"] == "NÃO").sum())
        resumo.append({
            "unidade": unidade,
            "prioritarios": p,
            "nao_prioritarios": n,
            "total": int(len(bloco)),
        })

    return {
        "total_analise_tecnica": int(len(df)),
        "fora_distribuicao": int(fora.sum()),
        "problemas": problemas,
        "resumo": resumo,
    }


def gerar_acompanhamento_zip(conteudo: bytes) -> tuple[bytes, dict]:
    df = preparar_datas(ler_demandas(conteudo))
    df = df[
        _normalizar_texto(df[COL_STATUS]).str.upper().eq("ANÁLISE TÉCNICA")
    ].copy()

    df["PRIORIDADE_NORMALIZADA"] = _normalizar_texto(df[COL_PRIORIDADE]).str.upper()
    df["UNIDADE_ACOMPANHAMENTO"] = df[COL_AREA].apply(identificar_unidade_acompanhamento)

    fora_mask = df["UNIDADE_ACOMPANHAMENTO"].isna()
    aptos = df[
        df["UNIDADE_ACOMPANHAMENTO"].notna()
        & df["PRIORIDADE_NORMALIZADA"].isin(["SIM", "NÃO"])
    ].copy()

    unidades = ["SA", "SD", "SF", "SH", "SR"]
    conselhos = sorted(
        u for u in aptos["UNIDADE_ACOMPANHAMENTO"].dropna().unique()
        if str(u).startswith("CD-DIRETOR-")
    )
    unidades.extend(conselhos)

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for unidade in unidades:
            bloco = aptos[aptos["UNIDADE_ACOMPANHAMENTO"] == unidade].copy()
            if bloco.empty:
                continue
            zf.writestr(f"{unidade}.xlsx", _xlsx_unidade(bloco))

        if fora_mask.any():
            zf.writestr("FORA_DA_DISTRIBUICAO.xlsx", _xlsx_fora(df[fora_mask].copy()))

    analise = analisar_acompanhamento(conteudo)
    nome = f"acompanhamento_superintendencias_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    return zip_buffer.getvalue(), {"nome": nome, **analise}
